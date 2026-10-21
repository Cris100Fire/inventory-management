import csv
from io import BytesIO
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from django.utils import timezone

STOCK_REPORT_HEADERS = [
    "Producto",
    "SKU",
    "Categoría",
    "Proveedor",
    "Precio",
    "Stock",
    "Stock mínimo",
    "Estado",
    "Valor estimado",
]

MOVEMENT_REPORT_HEADERS = [
    "Fecha",
    "Producto",
    "SKU",
    "Tipo de movimiento",
    "Cantidad",
    "Usuario",
    "Proveedor",
    "Observaciones",
]


def safe_spreadsheet_text(value):
    text = str(value)

    if text.startswith(("=", "+", "-", "@", "\t", "\r", "\n")):
        return "'" + text

    return text


def get_stock_status(product):
    if product.stock == 0:
        return "Agotado"

    if product.stock <= product.minimum_stock:
        return "Stock bajo"

    return "Disponible"


def get_stock_report_rows(queryset):
    for product in queryset.iterator():
        yield [
            safe_spreadsheet_text(product.name),
            safe_spreadsheet_text(product.sku),
            safe_spreadsheet_text(product.category.name),
            safe_spreadsheet_text(
                product.supplier.name if product.supplier else "Sin proveedor"
            ),
            product.price,
            product.stock,
            product.minimum_stock,
            get_stock_status(product),
            product.stock_value,
        ]


def export_stock_csv(queryset, filename="reporte_existencias.csv"):
    response = HttpResponse(
        content_type="text/csv; charset=utf-8",
    )

    response["Content-Disposition"] = f'attachment; filename="{filename}"'

    # BOM para mejorar la compatibilidad con Excel.
    response.write("\ufeff")

    writer = csv.writer(response)

    # Encabezados
    writer.writerow(STOCK_REPORT_HEADERS)

    # Datos
    for row in get_stock_report_rows(queryset):
        writer.writerow(row)

    return response


def export_stock_excel(
    queryset, filename="reporte_existencias.xlsx", sheet_name="Existencias"
):
    workbook = Workbook()

    worksheet = workbook.active

    # Pylance considera que workbook.active puede ser None.
    # Esta comprobación garantiza que la hoja existe.
    assert worksheet is not None

    worksheet.title = sheet_name

    # =========================================
    # ENCABEZADOS
    # =========================================

    worksheet.append(STOCK_REPORT_HEADERS)

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="2563EB",
    )

    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    worksheet.row_dimensions[1].height = 28

    # =========================================
    # DATOS
    # =========================================

    for row in get_stock_report_rows(queryset):
        worksheet.append(row)

    # =========================================
    # FORMATOS NUMÉRICOS
    # =========================================

    for row in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row,
    ):
        # Precio
        row[4].number_format = "#,##0.00"

        # Stock
        row[5].number_format = "#,##0"

        # Stock mínimo
        row[6].number_format = "#,##0"

        # Valor estimado
        row[8].number_format = "#,##0.00"

    # =========================================
    # ANCHOS DE COLUMNAS
    # =========================================

    column_widths = [
        30,  # Producto
        18,  # SKU
        22,  # Categoría
        26,  # Proveedor
        16,  # Precio
        12,  # Stock
        16,  # Stock mínimo
        18,  # Estado
        20,  # Valor estimado
    ]

    for index, width in enumerate(column_widths, start=1):
        column_letter = get_column_letter(index)

        worksheet.column_dimensions[column_letter].width = width

    # =========================================
    # CONFIGURACIÓN DE LA HOJA
    # =========================================

    # Mantener visibles los encabezados al desplazarse.
    worksheet.freeze_panes = "A2"

    # Habilitar filtros en Excel.
    worksheet.auto_filter.ref = worksheet.dimensions

    # =========================================
    # GENERAR ARCHIVO
    # =========================================

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument." "spreadsheetml.sheet"
        ),
    )

    response["Content-Disposition"] = f'attachment; filename="{filename}"'

    return response


def get_movement_report_rows(queryset):
    for movement in queryset.iterator():

        created_at = movement.created_at

        if timezone.is_aware(created_at):
            created_at = timezone.localtime(created_at)

        formatted_date = created_at.strftime("%d/%m/%Y %H:%M")

        yield [
            formatted_date,
            safe_spreadsheet_text(movement.product.name),
            safe_spreadsheet_text(movement.product.sku),
            movement.get_movement_type_display(),
            movement.quantity,
            safe_spreadsheet_text(movement.user.get_username()),
            safe_spreadsheet_text(
                movement.supplier.name if movement.supplier else "Sin proveedor"
            ),
            safe_spreadsheet_text(movement.notes or ""),
        ]


def export_movement_csv(queryset):
    response = HttpResponse(content_type="text/csv; charset=utf-8")

    response["Content-Disposition"] = 'attachment; filename="reporte_movimientos.csv"'

    # BOM para compatibilidad con Excel.
    response.write("\ufeff")

    writer = csv.writer(response)

    writer.writerow(MOVEMENT_REPORT_HEADERS)

    for row in get_movement_report_rows(queryset):
        writer.writerow(row)

    return response


def export_movement_excel(queryset):
    workbook = Workbook()

    worksheet = workbook.active

    # Evita advertencias de Pylance.
    assert worksheet is not None

    worksheet.title = "Movimientos"

    # Encabezados.
    worksheet.append(MOVEMENT_REPORT_HEADERS)

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="2563EB",
    )

    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    worksheet.row_dimensions[1].height = 28

    # Registros.
    for row in get_movement_report_rows(queryset):
        worksheet.append(row)

    # Formato de cantidades.
    for row in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row,
    ):
        row[4].number_format = "#,##0"

    # Ancho de columnas.
    column_widths = [
        20,  # Fecha
        32,  # Producto
        18,  # SKU
        24,  # Tipo
        14,  # Cantidad
        22,  # Usuario
        28,  # Proveedor
        45,  # Observaciones
    ]

    for index, width in enumerate(
        column_widths,
        start=1,
    ):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = worksheet.dimensions

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument." "spreadsheetml.sheet"
        ),
    )

    response["Content-Disposition"] = 'attachment; filename="reporte_movimientos.xlsx"'

    return response
