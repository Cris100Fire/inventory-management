import csv
from decimal import Decimal
from io import BytesIO, StringIO

from django.contrib.auth import get_user_model
from django.test import TestCase

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from apps.catalog.models import Category, Product
from apps.inventory.models import StockMovement

from apps.reports.services.queries import (
    get_stock_report_queryset,
    get_low_stock_report_queryset,
    get_movement_report_queryset,
)

from apps.reports.services.exporters import (
    STOCK_REPORT_HEADERS,
    MOVEMENT_REPORT_HEADERS,
    export_stock_csv,
    export_stock_excel,
    export_movement_csv,
    export_movement_excel,
    safe_spreadsheet_text,
)


class ReportExportTests(TestCase):

    @classmethod
    def setUpTestData(cls):

        cls.category = Category.objects.create(
            name="Categoría exportaciones",
        )

        cls.available = Product.objects.create(
            name="Teclado disponible",
            sku="EXP-001",
            price=Decimal("50.00"),
            stock=20,
            minimum_stock=5,
            category=cls.category,
        )

        cls.low_stock = Product.objects.create(
            name="Mouse stock bajo",
            sku="EXP-002",
            price=Decimal("25.00"),
            stock=3,
            minimum_stock=5,
            category=cls.category,
        )

        cls.out_of_stock = Product.objects.create(
            name="Monitor agotado",
            sku="EXP-003",
            price=Decimal("400.00"),
            stock=0,
            minimum_stock=5,
            category=cls.category,
        )

        User = get_user_model()

        cls.user = User.objects.create_user(
            username="export_user",
            password="TestPassword123!",
        )

        cls.entry = StockMovement.objects.create(
            product=cls.available,
            movement_type=StockMovement.MovementType.ENTRY,
            quantity=10,
            user=cls.user,
            notes="Ingreso de productos",
        )

        cls.exit = StockMovement.objects.create(
            product=cls.low_stock,
            movement_type=StockMovement.MovementType.EXIT,
            quantity=2,
            user=cls.user,
            notes="Salida de productos",
        )

    # ==========================================
    # UTILIDADES
    # ==========================================

    def read_csv(self, response):
        """
        Convierte una respuesta CSV en una lista de filas.
        """

        content = response.content.decode("utf-8-sig")

        return list(csv.reader(StringIO(content)))

    def read_excel(self, response):
        """
        Abre un archivo Excel generado en memoria.
        """

        workbook = load_workbook(
            BytesIO(response.content),
            data_only=True,
        )

        return workbook

    def get_active_worksheet(self, workbook) -> Worksheet:
        """
        Obtiene la hoja activa del libro Excel.

        Si no existe, la prueba falla con un
        mensaje descriptivo.
        """

        worksheet = workbook.active

        if worksheet is None:
            self.fail("El archivo Excel no contiene una hoja activa.")

        return worksheet

    # ==========================================
    # EXISTENCIAS - CSV
    # ==========================================

    def test_stock_csv_headers(self):

        queryset = get_stock_report_queryset()

        response = export_stock_csv(queryset)

        rows = self.read_csv(response)

        self.assertEqual(
            rows[0],
            STOCK_REPORT_HEADERS,
        )

    def test_stock_csv_contains_products(self):

        queryset = get_stock_report_queryset()

        response = export_stock_csv(queryset)

        rows = self.read_csv(response)

        product_names = [row[0] for row in rows[1:]]

        self.assertIn(
            "Teclado disponible",
            product_names,
        )

        self.assertIn(
            "Mouse stock bajo",
            product_names,
        )

        self.assertIn(
            "Monitor agotado",
            product_names,
        )

    def test_stock_csv_respects_filters(self):

        queryset = get_stock_report_queryset(
            {
                "stock_status": "out",
            }
        )

        response = export_stock_csv(queryset)

        rows = self.read_csv(response)

        self.assertEqual(
            len(rows),
            2,
        )

        self.assertEqual(
            rows[1][0],
            "Monitor agotado",
        )

    # ==========================================
    # EXISTENCIAS - EXCEL
    # ==========================================

    def test_stock_excel_sheet_name(self):

        queryset = get_stock_report_queryset()

        response = export_stock_excel(queryset)

        workbook = self.read_excel(response)

        worksheet = self.get_active_worksheet(workbook)

        self.assertEqual(
            worksheet.title,
            "Existencias",
        )

    def test_stock_excel_headers(self):

        queryset = get_stock_report_queryset()

        response = export_stock_excel(queryset)

        workbook = self.read_excel(response)

        worksheet = self.get_active_worksheet(workbook)

        headers = [cell.value for cell in worksheet[1]]

        self.assertEqual(
            headers,
            STOCK_REPORT_HEADERS,
        )

    def test_stock_excel_numeric_values(self):

        queryset = get_stock_report_queryset(
            {
                "search": "Teclado",
            }
        )

        response = export_stock_excel(queryset)

        workbook = self.read_excel(response)

        worksheet = self.get_active_worksheet(workbook)

        self.assertEqual(
            worksheet["F2"].value,
            20,
        )

        self.assertEqual(
            worksheet["I2"].value,
            1000,
        )

        self.assertIsInstance(
            worksheet["F2"].value,
            (int, float),
        )

    # ==========================================
    # STOCK BAJO - CSV
    # ==========================================

    def test_low_stock_csv_only_critical_products(self):

        queryset = get_low_stock_report_queryset()

        response = export_stock_csv(
            queryset,
            filename="reporte_stock_bajo.csv",
        )

        rows = self.read_csv(response)

        product_names = [row[0] for row in rows[1:]]

        self.assertNotIn(
            "Teclado disponible",
            product_names,
        )

        self.assertIn(
            "Mouse stock bajo",
            product_names,
        )

        self.assertIn(
            "Monitor agotado",
            product_names,
        )

    def test_low_stock_csv_filename(self):

        queryset = get_low_stock_report_queryset()

        response = export_stock_csv(
            queryset,
            filename="reporte_stock_bajo.csv",
        )

        self.assertIn(
            "reporte_stock_bajo.csv",
            response["Content-Disposition"],
        )

    # ==========================================
    # STOCK BAJO - EXCEL
    # ==========================================

    def test_low_stock_excel_sheet_name(self):

        queryset = get_low_stock_report_queryset()

        response = export_stock_excel(
            queryset,
            filename="reporte_stock_bajo.xlsx",
            sheet_name="Stock bajo",
        )

        workbook = self.read_excel(response)

        worksheet = self.get_active_worksheet(workbook)

        self.assertEqual(
            worksheet.title,
            "Stock bajo",
        )

    def test_low_stock_excel_only_critical_products(self):

        queryset = get_low_stock_report_queryset()

        response = export_stock_excel(
            queryset,
            filename="reporte_stock_bajo.xlsx",
            sheet_name="Stock bajo",
        )

        workbook = self.read_excel(response)

        worksheet = self.get_active_worksheet(workbook)

        product_names = [
            row[0]
            for row in worksheet.iter_rows(
                min_row=2,
                values_only=True,
            )
        ]

        self.assertEqual(
            len(product_names),
            2,
        )

        self.assertNotIn(
            "Teclado disponible",
            product_names,
        )

        self.assertIn(
            "Mouse stock bajo",
            product_names,
        )

        self.assertIn(
            "Monitor agotado",
            product_names,
        )

    # ==========================================
    # MOVIMIENTOS - CSV
    # ==========================================

    def test_movement_csv_headers(self):

        queryset = get_movement_report_queryset()

        response = export_movement_csv(queryset)

        rows = self.read_csv(response)

        self.assertEqual(
            rows[0],
            MOVEMENT_REPORT_HEADERS,
        )

    def test_movement_csv_contains_movements(self):

        queryset = get_movement_report_queryset()

        response = export_movement_csv(queryset)

        rows = self.read_csv(response)

        self.assertEqual(
            len(rows),
            3,
        )

        movement_types = [row[3] for row in rows[1:]]

        self.assertIn(
            "Entrada",
            movement_types,
        )

        self.assertIn(
            "Salida",
            movement_types,
        )

    def test_movement_csv_respects_filters(self):

        queryset = get_movement_report_queryset(
            {
                "movement_type": StockMovement.MovementType.ENTRY,
            }
        )

        response = export_movement_csv(queryset)

        rows = self.read_csv(response)

        self.assertEqual(
            len(rows),
            2,
        )

        self.assertEqual(
            rows[1][3],
            "Entrada",
        )

        self.assertEqual(
            rows[1][4],
            "10",
        )

    # ==========================================
    # MOVIMIENTOS - EXCEL
    # ==========================================

    def test_movement_excel_sheet_name(self):

        queryset = get_movement_report_queryset()

        response = export_movement_excel(queryset)

        workbook = self.read_excel(response)

        worksheet = self.get_active_worksheet(workbook)

        self.assertEqual(
            worksheet.title,
            "Movimientos",
        )

    def test_movement_excel_headers(self):

        queryset = get_movement_report_queryset()

        response = export_movement_excel(queryset)

        workbook = self.read_excel(response)

        worksheet = self.get_active_worksheet(workbook)

        headers = [cell.value for cell in worksheet[1]]

        self.assertEqual(
            headers,
            MOVEMENT_REPORT_HEADERS,
        )

    def test_movement_excel_quantity_is_numeric(self):

        queryset = get_movement_report_queryset(
            {
                "movement_type": StockMovement.MovementType.ENTRY,
            }
        )

        response = export_movement_excel(queryset)

        workbook = self.read_excel(response)

        worksheet = self.get_active_worksheet(workbook)

        self.assertEqual(
            worksheet["E2"].value,
            10,
        )

        self.assertIsInstance(
            worksheet["E2"].value,
            (int, float),
        )

    # ==========================================
    # SEGURIDAD DE EXPORTACIONES
    # ==========================================

    def test_spreadsheet_formula_protection(self):

        dangerous_values = [
            "=SUM(1,1)",
            "+123",
            "-123",
            "@formula",
            "\tformula",
        ]

        for value in dangerous_values:

            with self.subTest(value=value):

                safe_value = safe_spreadsheet_text(value)

                self.assertTrue(safe_value.startswith("'"))

    def test_normal_text_is_not_modified(self):

        value = "Producto normal"

        self.assertEqual(
            safe_spreadsheet_text(value),
            value,
        )
